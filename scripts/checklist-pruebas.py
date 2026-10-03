# -*- coding: utf-8 -*-
"""Genera el checklist de pruebas de la Plataforma IWL.

    pip3 install openpyxl
    python3 scripts/checklist-pruebas.py

Escribe IWL_Checklist_Pruebas.xlsx en la raíz, SOBRESCRIBIENDO lo que haya.
Si alguien está rellenando una ronda, que se guarde una copia antes.

Las pruebas viven aquí y no en el .xlsx a propósito: así se añaden en el
mismo commit que la función que hay que probar, y se revisan como código.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

MAGENTA = "FF007A"
CARBON = "0D0D11"
ZINC = "3F3F46"
PAPEL = "FAFAFA"
FILETE = "E4E4E7"

wb = Workbook()

# ---------------------------------------------------------------- Pruebas
ws = wb.active
ws.title = "Pruebas"

CABECERAS = [
    ("Nº", 9),
    ("Módulo", 24),
    ("Pantalla", 30),
    ("Entra como", 22),
    ("Qué se prueba", 46),
    ("Pasos", 62),
    ("Qué debería pasar", 62),
    ("Resultado", 14),
    ("Gravedad", 14),
    ("Comentarios", 50),
]

# (módulo, pantalla, rol, qué, pasos, esperado)
P = []

def add(mod, pant, rol, que, pasos, esperado):
    P.append((mod, pant, rol, que, pasos, esperado))

# --- Acceso y cuentas
M = "1. Acceso y cuentas"
add(M, "/entrar", "admin@iwl.test", "Entrar con correo y contraseña",
    "Escribir correo y contraseña y pulsar Entrar.",
    "Entra y aterriza en la cartera.")
add(M, "/entrar", "cualquiera", "Contraseña equivocada",
    "Escribir una contraseña que no es y pulsar Entrar.",
    "Dice qué ha pasado y qué hacer, sin signos de exclamación. No se queda colgado.")
add(M, "/entrar", "—", "No hay autoservicio",
    "Buscar en la pantalla de entrada alguna forma de crear una cuenta.",
    "No la hay. Las cuentas las crea la dirección de IWL.")
add(M, "/", "fundadora@marea.test", "La raíz lleva a donde toca según quién eres",
    "Entrar y abrir la dirección raíz.",
    "Una fundadora cae en su proyecto; IWL, en la cartera.")
add(M, "/cartera", "fundadora@marea.test", "Una fundadora no entra en la cartera",
    "Estando dentro como fundadora, escribir /cartera en la barra de direcciones.",
    "La devuelve a su proyecto. No ve la cartera ni de refilón.")
add(M, "/admin", "programa@iwl.test", "El equipo de IWL no configura el programa",
    "Entrar como equipo_iwl y escribir /admin en la barra.",
    "No entra. La configuración es solo de la dirección.")
add(M, "/perfil", "fundadora@marea.test", "Cambiar mi propia contraseña",
    "Abrir Mi cuenta, poner contraseña nueva y guardar. Salir y volver a entrar con ella.",
    "Se guarda y la nueva funciona.")
add(M, "/sin-compania", "una cuenta recién creada", "Quien no tiene compañía lo entiende",
    "Entrar con una cuenta recién creada y sin asignar.",
    "Dice qué pasa y a quién pedírselo, no una pantalla en blanco.")
add(M, "Cualquiera", "—", "Salir",
    "Pulsar Salir desde cualquier pantalla.",
    "Vuelve a /entrar y la sesión queda cerrada de verdad (atrás no reabre nada).")

# --- Embudo: convocatoria
M = "2. Convocatoria"
add(M, "/embudo", "admin@iwl.test", "Abrir la convocatoria de una cohorte",
    "Embudo → Convocatoria → elegir cohorte, marcar abierta, guardar.",
    "Confirma. Arriba aparece el aviso de que esa convocatoria está abierta.")
add(M, "/embudo", "admin@iwl.test", "La dirección que se reparte está a la vista",
    "Con la convocatoria abierta, mirar el bloque «Reparte esta dirección» y pulsar Copiar.",
    "Se ve la dirección entera y se copia al portapapeles.")
add(M, "/presentarse", "sin cuenta", "Con la convocatoria abierta, el formulario sale",
    "Abrir la dirección copiada en una ventana privada.",
    "Sale el formulario de candidatura, sin pedir cuenta.")
add(M, "/embudo", "admin@iwl.test", "Cerrar la convocatoria",
    "Desmarcar abierta y guardar. Volver a abrir /presentarse en ventana privada.",
    "Dice que no hay convocatoria en curso Y da una dirección de correo para avisar. Antes decía «escríbenos» sin decir dónde.")
add(M, "/embudo", "admin@iwl.test", "Fecha de cierre",
    "Abrir la convocatoria con una fecha de cierre y mirar el texto de arriba.",
    "Dice hasta cuándo está abierta.")

# --- Presentarse
M = "3. Presentar candidatura"
add(M, "/presentarse", "sin cuenta", "Presentarse con todo relleno",
    "Rellenar los campos, dejar dos enlaces de Drive y enviar.",
    "«Recibida. Gracias.» y debajo el enlace privado con botón de copiar.")
add(M, "/presentarse", "sin cuenta", "Solo lo obligatorio",
    "Rellenar lo que lleva asterisco: startup, en qué punto estáis, tu nombre y tu correo. Enviar.",
    "Se acepta. El resto es opcional de verdad.")
add(M, "/presentarse", "sin cuenta", "La etapa no se puede esquivar",
    "Mirar el desplegable «¿En qué punto estáis?» e intentar enviar sin tocarlo.",
    "Ya no existe «Prefiero no decirlo». Pone «Elige una» y no deja enviar sin elegir.")
add(M, "/presentarse", "sin cuenta", "Un porcentaje escrito a mano (el caso dronesec)",
    "En «Liderazgo femenino (%)» escribir «50 %». También probar «33,3».",
    "Se acepta. Llega como número, no se rechaza por el signo ni por la coma.")
add(M, "/presentarse", "sin cuenta", "Cuando algo falla, se dice cuál y no se borra nada",
    "Rellenarlo entero y poner «la mitad» en el porcentaje. Enviar.",
    "Dice «Revisa este campo: liderazgo femenino», lo marca en rojo, y TODO lo demás sigue escrito.")
add(M, "/presentarse", "sin cuenta", "Correo mal escrito",
    "Poner un correo sin arroba y enviar.",
    "Lo señala sin perder el resto de lo escrito.")
add(M, "/candidatura/<testigo>", "sin cuenta", "El enlace privado abre su candidatura",
    "Pegar el enlace que dio el formulario en otra ventana.",
    "Ve su nombre, el estado «Recibida» y lo que mandó.")
add(M, "/candidatura/<testigo>", "sin cuenta", "Añadir material que se olvidó",
    "En su enlace, Añadir un documento: título y URL. Añadir.",
    "«Añadido. Gracias.» y el documento aparece en la lista.")
add(M, "/candidatura/xxxx", "sin cuenta", "Un enlace inventado no dice si existe",
    "Escribir un testigo cualquiera inventado.",
    "404 a secas. No distingue entre «no existe» y «no es tuyo».")

# --- Alta a mano
M = "4. Alta a mano"
add(M, "/embudo", "admin@iwl.test", "Dar de alta una candidatura conocida en un evento",
    "Embudo → Dar de alta una candidatura a mano: cohorte, nombre, contacto, correo.",
    "Confirma que entra en el embudo y aparece en la lista de abiertas.")
add(M, "/embudo/<id>", "admin@iwl.test", "El enlace de la invitada se ve para mandarlo",
    "Abrir su ficha en el embudo y buscar el enlace privado.",
    "Está a la vista, se puede copiar y mandar a mano.")
add(M, "/candidatura/<testigo>", "sin cuenta", "A la invitada se le pide que se presente",
    "Abrir el enlace de una candidatura dada de alta por IWL.",
    "Dice «Cuéntanos quiénes sois», NO «Recibida, gracias». Y tiene formulario para contarlo.")
add(M, "/candidatura/<testigo>", "sin cuenta", "La invitada completa su propia ficha",
    "Rellenar descripción, punto en que están, web, equipo. Guardar.",
    "«Guardado. Gracias.» y lo escrito queda visible.")
add(M, "/candidatura/<testigo>", "sin cuenta", "Un formulario a medias no borra lo que ya había",
    "IWL apunta el sector al dar de alta; ella deja ese campo vacío y guarda.",
    "El sector se queda como estaba. No se vacía.")

# --- Recorrido del embudo
M = "5. Recorrido del embudo"
add(M, "/embudo", "admin@iwl.test", "La foto del embudo cuadra",
    "Mirar las cuatro cifras de arriba y el recorrido paso a paso.",
    "En proceso, firmadas, descartadas y conversión cuadran con las listas de abajo.")
add(M, "/embudo", "admin@iwl.test", "La conversión no miente con pocos datos",
    "Mirar la conversión con menos de cinco candidaturas cerradas.",
    "Pone «—» y explica que hacen falta cinco. No inventa un porcentaje.")
add(M, "/embudo/<id>", "admin@iwl.test", "Mover de paso",
    "Pasar una candidatura de presentada a en revisión, y luego a reunión.",
    "Cambia el paso y queda anotado en el historial con fecha.")
add(M, "/embudo/<id>", "admin@iwl.test", "Anotar lo que pasó en la reunión",
    "Anotar un evento con lo hablado en la reunión y en el comité.",
    "Queda en el historial, lo más reciente arriba.")
add(M, "/embudo/<id>", "admin@iwl.test", "Preseleccionar",
    "Mover a preseleccionada después del comité.",
    "Cambia de paso y la cuenta del recorrido se mueve con ella.")
add(M, "/embudo/<id>", "admin@iwl.test", "Descartar con su motivo",
    "Descartar una candidatura escribiendo el motivo.",
    "Sale de las abiertas, entra en descartadas, y queda apuntado EN QUÉ PASO se cayó.")
add(M, "/embudo", "admin@iwl.test", "Dónde se cae la gente",
    "Mirar el bloque de descartadas y las pastillas por paso.",
    "Dice cuántas se cayeron en cada paso. Es la pregunta por la que se cambia un proceso.")
add(M, "/embudo/<id>", "admin@iwl.test", "Reabrir una descartada",
    "Reabrir una que se descartó por error.",
    "Vuelve a las abiertas, al paso donde estaba.")
add(M, "/candidatura/<testigo>", "sin cuenta", "La descartada NO lee el motivo",
    "Abrir el enlace de una candidatura descartada.",
    "Dice «Proceso cerrado» y que se lo contarán. El motivo NO aparece, y tampoco la caja de documentos vacía reprochándole lo que no mandó.")
add(M, "/candidatura/<testigo>", "sin cuenta", "La candidata ve momentos, no jerga",
    "Mirar el enlace de una que está en comité.",
    "Dice «En estudio». No dice «comité» ni el paso exacto, y NO le pide que se presente: ya hemos hablado con ella.")
add(M, "/candidatura/<testigo>", "sin cuenta", "Corregir un dato estando ya avanzada",
    "En una candidatura en comité, buscar cómo cambiar la web o el teléfono.",
    "Está plegado bajo «Corregir nuestros datos». Existe, pero no es lo primero que ve.")

# --- NDA y due diligence de la candidata
M = "6. NDA y DD de candidata"
add(M, "/embudo/<id>", "admin@iwl.test", "Al llegar al NDA se piden los documentos",
    "Mover una preseleccionada al paso de NDA.",
    "Aparece sola la lista de documentos que hay que pedirle para el due diligence.")
add(M, "/embudo/<id>", "admin@iwl.test", "Si la candidatura retrocede, los deberes desaparecen",
    "Con una en due diligence, devolverla a comité y mirar su enlace.",
    "Ya no le pide documentos. Lo que entregó NO se borra: vuelve a salir si se la devuelve a diligencia.")
add(M, "/embudo/<id>", "admin@iwl.test", "Pedir un documento concreto de más",
    "Añadir una petición con su nombre y para cuándo.",
    "Se añade a la lista y ella la ve en su pantalla.")
add(M, "/embudo/<id>", "admin@iwl.test", "Darle cuenta para el due diligence",
    "Pulsar «Darle cuenta para el due diligence».",
    "Enseña el correo y la contraseña UNA SOLA VEZ, y avisa de que no se vuelve a enseñar. Copiarla.")
add(M, "/candidatura/<testigo>", "sin cuenta", "El enlace privado deja de valer",
    "Después de darle cuenta, abrir el enlace privado de antes.",
    "404. Desde que tiene cuenta, se entra por la puerta.")
add(M, "/entrar → /candidatura", "la candidata", "Entra con su cuenta y aterriza en lo suyo",
    "Entrar con el correo y la contraseña que se enseñaron.",
    "Cae en /candidatura, con su nombre y lo que le piden. NO en «no tienes compañía».")
add(M, "/candidatura", "la candidata", "Sube un documento del due diligence",
    "Elegir una petición y subir el fichero.",
    "Queda entregado y ella lo ve. El fichero es privado: no hay URL pública.")
add(M, "/embudo/<id>", "admin@iwl.test", "IWL abre lo que entregó",
    "Abrir desde la ficha el documento que subió.",
    "Se descarga o se abre con enlace firmado y caducable.")

# --- Acuerdo y firma
M = "7. Acuerdo y firma"
add(M, "/embudo/<id>", "admin@iwl.test", "Fijar la etapa tras el due diligence",
    "En el bloque «El acuerdo», el campo «Etapa, tras el due diligence».",
    "Se guarda y AL RECARGAR sigue puesta: el desplegable ya no vuelve a donde estaba. Es la etapa que hereda la compañía al firmar.")
add(M, "/embudo/<id>", "admin@iwl.test", "Armar el acuerdo con el equity",
    "Escribir lo que aporta IWL y el equity que se pide. Guardar.",
    "Queda guardado y visible en el bloque de NDA, estado y equity.")
add(M, "/embudo/<id>", "admin@iwl.test", "Firmar el acuerdo",
    "Pulsar firmar con la fecha.",
    "Pasa a firmada, SE CREA LA COMPAÑÍA, y aparece el enlace «Ver su ficha».")
add(M, "/cartera/<slug>", "admin@iwl.test", "La compañía nace con lo del embudo",
    "Abrir la ficha recién creada desde el enlace.",
    "Trae nombre, sector, descripción y etapa. No hay que volver a teclear nada.")
add(M, "/proyecto/diligencia", "la fundadora", "La compañía lee su expediente de candidatura",
    "Entrar como la fundadora de la compañía recién creada y buscar sus documentos.",
    "Los ve. NO están duplicados: se leen del expediente, no se copiaron.")
add(M, "/embudo/<id>", "admin@iwl.test", "Una firmada no se edita",
    "Intentar cambiar el equity de una ya firmada, y guardar.",
    "Lo rechaza diciendo que ya está firmada. Lo impide la base, no la pantalla: forzarlo tampoco vale.")

# --- Cartera
M = "8. Cartera y cohorte"
add(M, "/cartera", "programa@iwl.test", "La cohorte de un vistazo",
    "Abrir la cartera y recorrer la pantalla de arriba abajo.",
    "Cifras de cabecera, cohorte, scorecards y siguientes pasos, sin texto cortado ni solapes.")
add(M, "/cartera", "admin@iwl.test", "Subir el logo de una startup",
    "Al lado de una compañía, subir una imagen de logo.",
    "Se ve el logo en su sitio y la confirmación NO desaparece sola.")
add(M, "/cartera", "admin@iwl.test", "Quitar un logo",
    "Quitar el logo de una compañía.",
    "Vuelve a las iniciales y lo confirma.")
add(M, "/cartera", "programa@iwl.test", "Preparación en el tiempo",
    "Mirar el gráfico y su texto de abajo.",
    "Se lee entero, sin recortes. Cada serie lleva etiqueta de texto, no solo color.")
add(M, "/cartera", "programa@iwl.test", "Siguientes pasos llevan a donde se hacen",
    "Pulsar un paso de «Siguientes pasos», en la cartera y en la ficha de una compañía.",
    "Abre la sección donde se actúa: un hallazgo lleva al DD técnico, un hito del Anexo a Programa, el runway a KPI.")
add(M, "/cartera", "programa@iwl.test", "Una compañía archivada no sale",
    "Archivar una desde Administración y volver a la cartera.",
    "Deja de aparecer en la cohorte. No se ha borrado nada.")
add(M, "/cartera/<slug>", "programa@iwl.test", "Entrar en una compañía",
    "Pulsar sobre una compañía de la lista.",
    "Abre su resumen con las ocho secciones arriba.")

# --- Comparativa
M = "9. Comparativa"
add(M, "/comparativa", "programa@iwl.test", "¿Quién va por delante?",
    "Mirar el primer gráfico.",
    "Barras con el % del objetivo de SU etapa alcanzado. 100 es llegar. No niveles brutos.")
add(M, "/comparativa", "programa@iwl.test", "¿Dónde flojea la cohorte?",
    "Mirar el segundo gráfico.",
    "Dimensiones ordenadas por brecha media, lo peor arriba, con cuántas compañías están por debajo.")
add(M, "/comparativa", "programa@iwl.test", "Las cifras exactas siguen estando",
    "Bajar al final de la página.",
    "La tabla «Las cifras, una a una» está ahí para cuando hace falta el dato exacto.")
add(M, "/comparativa", "programa@iwl.test", "Runway, MRR y preparación",
    "Mirar los tres bloques de tracción.",
    "Cuadran con lo cargado en KPI de cada compañía.")
add(M, "/comparativa", "programa@iwl.test", "Lo sin medir no es un cero",
    "Mirar una compañía sin evaluar todavía.",
    "Aparece como no evaluada («—»), NO como un cero ni la última del ranking.")

# --- Resumen de compañía
M = "10. Resumen de compañía"
add(M, "/cartera/<slug>", "programa@iwl.test", "La cabecera dice dónde está",
    "Mirar score técnico, preparación, semáforo, fase, etapa y perfil.",
    "Cuadran con lo evaluado. El semáforo explica su motivo.")
add(M, "/cartera/<slug>", "programa@iwl.test", "Movimiento frente a la línea base",
    "Mirar la variación al lado de cada score.",
    "Dice cuánto se ha movido desde que se congeló la línea base.")
add(M, "/cartera/<slug>", "programa@iwl.test", "Sin evaluar parece sin evaluar",
    "Abrir una compañía a la que nadie ha puntuado.",
    "Los scores ponen «—». No hay un 0 que parezca una mala nota.")
add(M, "/proyecto", "fundadora@marea.test", "La fundadora ve lo mismo en papel blanco",
    "Entrar como fundadora y abrir su proyecto.",
    "Mismos datos, tema claro, y solo lo suyo.")

# --- Hoja de ruta
M = "11. Hoja de ruta"
add(M, "/admin/rutas", "admin@iwl.test", "Crear una plantilla de hoja de ruta",
    "Crear plantilla, añadirle etapas e hitos.",
    "Se guarda y queda disponible para instanciar.")
add(M, "/cartera/<slug>/ruta", "admin@iwl.test", "Diseñar la hoja de ruta de una compañía",
    "Instanciar la plantilla sobre una compañía.",
    "Se COPIAN las etapas e hitos. Queda la hoja de ruta montada.")
add(M, "/admin/rutas", "admin@iwl.test", "Editar la plantilla no toca lo que está en marcha",
    "Cambiar un hito de la plantilla y volver a la hoja de ruta ya instanciada.",
    "La de la compañía NO cambia. Y al revés tampoco.")
add(M, "/cartera/<slug>/ruta", "admin@iwl.test", "Crear, editar y borrar una etapa",
    "Añadir una etapa, renombrarla y borrarla con «Borrar la etapa».",
    "Las tres funcionan. Al borrar avisa de qué pasa con sus hitos: no se borran, se quedan sueltos.")
add(M, "/cartera/<slug>/ruta", "admin@iwl.test", "Crear y mover un hito",
    "Añadir un hito con «Añadir un hito», y moverlo con el desplegable de su fila.",
    "Cambia de carril y conserva lo demás. Un hito suelto también se puede colocar en una etapa.")
add(M, "/cartera/<slug>/ruta", "admin@iwl.test", "Fijar el estado de entrada",
    "Marcar en qué punto entró la compañía.",
    "Queda fijado y sirve de referencia para el avance.")
add(M, "/proyecto/ruta", "fundadora@marea.test", "La fundadora registra un avance",
    "Registrar un avance sobre un hito.",
    "Cae en su carril, con su fecha, y se ve desde IWL.")
add(M, "/proyecto/ruta", "fundadora@marea.test", "Borrar un avance mal metido",
    "Borrar el avance que se acaba de registrar.",
    "Desaparece y el hito vuelve a su estado anterior.")
add(M, "/cartera/<slug>/ruta", "programa@iwl.test", "Un hito que condiciona el estado invertible",
    "Marcar un hito como condicionante y dejarlo sin cumplir.",
    "El semáforo de invertible lo refleja y dice por qué.")

# --- Programa y Anexo
M = "12. Programa y Anexo"
add(M, "/cartera/<slug>/programa", "admin@iwl.test", "Abrir el Anexo donde no hay ninguno",
    "En una compañía recién creada: Programa y Anexo → «Abrir el Anexo».",
    "Aparece el Anexo en borrador con sus campos. Antes aquí no había por dónde empezar y eso bloqueaba horas y aportación.")
add(M, "/cartera/<slug>/programa", "admin@iwl.test", "Rellenar el Anexo",
    "Poner duración, horas comprometidas y equity. Probar a escribir «7,5 %» en el equity.",
    "Se guarda, y el porcentaje con coma y signo se admite.")
add(M, "/cartera/<slug>/programa", "admin@iwl.test", "No se firma un Anexo vacío",
    "Abrir otro Anexo y pulsar firmar sin rellenar nada.",
    "No deja: pide al menos la duración y las horas. Lo firmado ya no se edita.")
add(M, "/cartera/<slug>/programa", "admin@iwl.test", "Firmar el Anexo",
    "Con la duración y las horas puestas, firmar.",
    "Pasa a «Firmado el …» y el formulario de edición desaparece. A partir de ahí se abre otra versión, no se edita.")
add(M, "/cartera/<slug>/programa", "admin@iwl.test", "Los hitos ya acordados del Anexo",
    "Mirar el bloque de hitos del Anexo y cambiar el estado de uno.",
    "Cambia y queda registrado.")
add(M, "/cartera/<slug>/programa", "mentor@iwl.test", "Registrar horas",
    "Imputar horas a una partida del Anexo.",
    "Se guardan con la tarifa COPIADA en la línea, no referenciada.")
add(M, "/admin/programa", "admin@iwl.test", "Cambiar una tarifa no reescribe el histórico",
    "Cambiar la tarifa de un perfil y volver al extracto de horas de antes.",
    "Las líneas viejas conservan la tarifa que tenían. Solo las nuevas usan la nueva.")
add(M, "/cartera/<slug>/programa", "programa@iwl.test", "Entregado sobre comprometido",
    "Mirar el compromiso del Anexo frente a lo entregado.",
    "Las dos cifras cuadran con las horas registradas.")
add(M, "/proyecto/programa", "fundadora@marea.test", "La fundadora objeta una imputación",
    "Objetar una línea de horas.",
    "Queda la objeción anotada y visible para IWL.")
add(M, "/cartera/<slug>/programa", "admin@iwl.test", "Un Anexo firmado no se edita",
    "Intentar cambiar una partida de un Anexo ya firmado.",
    "No deja. Lo firmado está congelado.")

# --- Aportación
M = "13. Aportación de IWL"
add(M, "/cartera/<slug>/aportacion", "mentor@iwl.test", "Registrar una aportación",
    "Registrar una sesión o una aportación con sus horas.",
    "Se guarda y suma en horas y valor.")
add(M, "/cartera/<slug>/aportacion", "programa@iwl.test", "Quién ha aportado",
    "Mirar el desglose por persona.",
    "Cuadra con lo registrado, con nombre y horas.")
add(M, "/proyecto/aportacion", "fundadora@marea.test", "La fundadora ve lo que ha recibido",
    "Abrir Aportación de IWL como fundadora.",
    "Ve horas y valor de mercado de lo que ha recibido.")

# --- DD técnico
M = "14. Due diligence técnico"
add(M, "/cartera/<slug>/tecnico", "revisor@niage.test", "Abrir una evaluación donde no había ninguna",
    "En una compañía recién creada: Due diligence técnico → «Abrir la evaluación técnica».",
    "Aparece el scorecard con el aviso «En borrador». Antes aquí no había por dónde empezar.")
add(M, "/proyecto/tecnico", "la fundadora", "La compañía NO ve el borrador",
    "Con la evaluación en borrador, mirar la sección como fundadora.",
    "Sigue diciendo que no hay evaluación publicada. El borrador es solo de quien evalúa.")
add(M, "/cartera/<slug>/tecnico", "revisor@niage.test", "Publicar la evaluación",
    "Puntuar al menos una dimensión, escribir el resumen y publicar.",
    "Se publica y la compañía ya la ve. Sin nada puntuado NO deja publicar.")
add(M, "/cartera/<slug>/tecnico", "revisor@niage.test", "Puntuar una dimensión",
    "Ajustar puntuación de una dimensión con su evidencia. Guardar.",
    "Se guarda y el score técnico se mueve.")
add(M, "/cartera/<slug>/tecnico", "revisor@niage.test", "Una puntuación sin evidencia no se guarda",
    "Poner un nivel y dejar la evidencia en una palabra. Guardar.",
    "No se guarda y lo dice: una puntuación sin evidencia no es una evaluación.")
add(M, "/proyecto/tecnico", "fundadora@marea.test", "Una fundadora no se puntúa a sí misma",
    "Buscar como fundadora alguna forma de ajustar la puntuación.",
    "No existe. Y si se forzara, la base lo rechaza. Es criterio de aceptación.")
add(M, "/cartera/<slug>/tecnico", "revisor@niage.test", "Registrar un hallazgo",
    "Registrar un hallazgo con su gravedad.",
    "Aparece en la lista de abiertos.")
add(M, "/cartera/<slug>/tecnico", "revisor@niage.test", "Cerrar un hallazgo",
    "Cambiar el estado de un hallazgo a cerrado.",
    "Sale de los abiertos y la cuenta de arriba baja.")
add(M, "/cartera/<slug>/tecnico", "revisor@niage.test", "Crear un punto de plan",
    "Crear un punto del plan de acción a partir de un hallazgo.",
    "Queda enlazado y con su estado.")
add(M, "/proyecto/tecnico", "fundadora@marea.test", "Responder el cuestionario",
    "Responder las preguntas del cuestionario técnico.",
    "Se guardan las respuestas y el revisor las ve.")
add(M, "/cartera/<slug>/tecnico", "programa@iwl.test", "El score se reparte sobre lo evaluado",
    "Evaluar solo tres dimensiones de las que haya y mirar el score.",
    "Se calcula sobre esas tres. Las no evaluadas NO cuentan como cero.")
add(M, "/cartera/<slug>/tecnico", "revisor2@niage.test", "Un revisor solo ve sus compañías",
    "Entrar como revisor2 y escribir el slug de una compañía que no lleva.",
    "No entra.")

# --- DD general
M = "15. Due diligence general"
add(M, "/cartera/<slug>/diligencia", "programa@iwl.test", "La checklist sale instanciada",
    "Abrir el due diligence general de una compañía.",
    "Tiene su propia copia de la checklist, no la plantilla compartida.")
add(M, "/cartera/<slug>/diligencia", "programa@iwl.test", "Cambiar el estado de un punto",
    "Marcar un punto como entregado o como no aplica.",
    "Cambia y la cuenta de pendientes se mueve.")
add(M, "/proyecto/diligencia", "fundadora@marea.test", "Subir desde la propia línea",
    "En la línea de un punto, pulsar «Subir» y elegir el fichero. Sin bajar al final de la página.",
    "Se sube ahí mismo, sin volver a elegir área ni punto, y el panel se cierra con Hecho.")
add(M, "/proyecto/diligencia", "fundadora@marea.test", "Al subir, el punto pasa a EN REVISIÓN",
    "Mirar el estado del punto justo después de subir.",
    "Dice «En revisión», no «Entregado». Y donde había un desplegable ahora hay una etiqueta.")
add(M, "/proyecto/diligencia", "fundadora@marea.test", "La compañía no saca un punto de revisión",
    "Intentar devolver a pendiente un punto que ya está en revisión.",
    "No puede: no hay control. Y si se forzara, la base lo rechaza.")
add(M, "/cartera/<slug>/diligencia", "revisor@niage.test", "IWL sí lo mueve",
    "Sobre ese mismo punto, pasarlo a validado.",
    "Cambia. El estado a partir de revisión es una valoración, y las valoraciones son de IWL.")
add(M, "/cartera/<slug>/diligencia", "programa@iwl.test", "Registrar un hallazgo general",
    "Registrar un hallazgo del due diligence general.",
    "Queda con su gravedad y se ve en la lista.")
add(M, "/cartera/<slug>/diligencia", "programa@iwl.test", "Un documento con caducidad",
    "Subir un documento poniéndole fecha de caducidad.",
    "Se guarda y avisa cuando esté caducado.")

# --- Business plan
M = "16. Business plan"
add(M, "/proyecto/plan", "fundadora@marea.test", "Adjuntar el business plan ya escrito",
    "Arriba del todo: «Adjuntar el business plan», elegir un PDF.",
    "Queda adjunto con su fecha y se abre con un clic. NO rellena las secciones: eso sigue siendo a mano.")
add(M, "/proyecto/plan", "fundadora@marea.test", "Subir otra versión reemplaza",
    "Adjuntar un segundo fichero.",
    "Sustituye al anterior. «El business plan» es singular, no una pila.")
add(M, "/proyecto/plan", "fundadora@marea.test", "Escribir una sección",
    "Rellenar una sección del business plan y guardar.",
    "Se guarda y se ve desde IWL.")
add(M, "/proyecto/plan", "fundadora@marea.test", "Hipótesis",
    "Añadir una hipótesis con cómo se va a validar.",
    "Queda registrada.")
add(M, "/cartera/<slug>/plan", "mentor@iwl.test", "Comentar una sección",
    "Dejar un comentario sobre una sección.",
    "La fundadora lo ve, con quién lo escribió.")
add(M, "/cartera/<slug>/plan", "admin@iwl.test", "La plantilla se copia al instanciar",
    "Cambiar la plantilla del business plan y mirar uno ya en marcha.",
    "El que está en marcha NO cambia.")

# --- KPI
M = "17. KPI y updates"
add(M, "/proyecto/kpi", "fundadora@marea.test", "Cargar los KPI del mes",
    "Cargar los valores del mes y guardar.",
    "Se guardan y aparecen en la serie.")
add(M, "/proyecto/kpi", "fundadora@marea.test", "Las derivadas se calculan solas",
    "Mirar runway, crecimiento y demás después de cargar.",
    "Salen calculadas. NO hay que teclearlas, y no se dejan editar a mano.")
add(M, "/proyecto/kpi", "fundadora@marea.test", "Escribir el update del mes",
    "Escribir logros, bloqueos y lo que viene. Guardar.",
    "Se guarda y se ve desde IWL con su fecha.")
add(M, "/cartera/<slug>/kpi", "programa@iwl.test", "La serie en el tiempo",
    "Mirar el gráfico de varios meses.",
    "Cuadra con lo cargado. Cada serie lleva etiqueta de texto.")
add(M, "/proyecto/kpi", "fundadora@marea.test", "Un mes sin cargar no es un cero",
    "Dejar un mes sin cargar y mirar el gráfico.",
    "Se ve el hueco. No se dibuja una caída a cero.")

# --- Línea base
M = "18. Línea base"
add(M, "/cartera/<slug>", "admin@iwl.test", "Congelar la línea base",
    "Congelar la línea base de una compañía.",
    "Queda congelada y a partir de ahí hay con qué comparar.")
add(M, "/cartera/<slug>", "admin@iwl.test", "Una línea base congelada no se edita",
    "Intentar cambiar una línea base ya congelada.",
    "No deja. Está congelada.")
add(M, "/cartera/<slug>", "programa@iwl.test", "La instantánea usa el mismo cálculo",
    "Comparar la cifra de la instantánea con la que se ve en pantalla.",
    "Son la misma. El cálculo no está reimplementado en dos sitios.")

# --- Administración
M = "19. Administración"
add(M, "/admin/companias", "admin@iwl.test", "Dar de alta una compañía a mano",
    "Crear una compañía con su cohorte y su hoja de ruta.",
    "Queda lista para trabajar, sin pasos sueltos pendientes.")
add(M, "/admin/companias", "admin@iwl.test", "Editar, archivar y restaurar",
    "Editar los datos, archivarla y restaurarla.",
    "Las tres funcionan. Archivar saca de la cartera y conserva el histórico entero.")
add(M, "/admin/companias", "admin@iwl.test", "Borrar una compañía vacía",
    "Dar de alta una y borrarla sin tocar nada más.",
    "Sale «Borrar» a secas. Pide escribir el identificador y se va.")
add(M, "/admin/companias", "admin@iwl.test", "Borrar una con histórico",
    "En una que tenga Anexo u horas: «Borrar con su histórico».",
    "Primero recuerda que lo normal es archivar. Luego dice QUÉ destruye («1 Anexo, 3 horas…») y pide el identificador exacto.")
add(M, "/admin/companias", "admin@iwl.test", "La lista se entera de lo que pasa en otra pantalla",
    "Abrir un Anexo de una compañía y volver a Compañías pulsando en el menú.",
    "Ya dice «Borrar con su histórico», no «Borrar» a secas. Si ofreciera el borrado simple, la base lo rechazaría.")
add(M, "/admin/personas", "admin@iwl.test", "Crear una persona con su contraseña",
    "Dar de alta a una persona. Dejar la contraseña en blanco para que la genere ella.",
    "Se enseña la contraseña generada —tres palabras y un número, dictable por teléfono— con aviso de que no se repite.")
add(M, "/admin/personas", "admin@iwl.test", "Restablecer la contraseña de alguien que se quedó fuera",
    "En la fila de esa persona: «Restablecer contraseña» → Generar → Guardar contraseña.",
    "El panel SE QUEDA ABIERTO con la contraseña a la vista y botón de copiar, hasta pulsar Hecho. Y entra con ella.")
add(M, "/admin/personas", "admin@iwl.test", "El verbo cambia según el caso",
    "Comparar la fila de alguien que nunca entró con la de alguien que sí.",
    "«Poner contraseña» en la primera, «Restablecer contraseña» en la segunda.")
add(M, "/admin/personas", "admin@iwl.test", "La propia NO se cambia desde el panel",
    "Buscar el botón de contraseña en tu propia fila.",
    "No está. En su lugar, «Tu contraseña se cambia en Mi cuenta», que lleva a /perfil. Desde el panel te echaría la sesión.")
add(M, "/admin/personas", "admin@iwl.test", "La lista se puede recorrer",
    "Mirar la lista entera de personas.",
    "Por cada una: nombre, posición, correo y dónde está asignada. Ni áreas ni biografía: eso está en su ficha.")
add(M, "/admin/personas/<id>", "admin@iwl.test", "La ficha de una persona",
    "Pulsar sobre el nombre de alguien.",
    "Se abre su ficha: proyectos, horas imputadas, tareas abiertas, áreas, descripción y asignaciones con detalle.")
add(M, "/admin/personas/<id>", "admin@iwl.test", "Adjuntar un CV",
    "En «Quién es», adjuntar un PDF y luego abrirlo.",
    "Se guarda y se abre con enlace firmado. Lo ve IWL y la propia persona; nadie más.")
add(M, "/admin/personas", "admin@iwl.test", "Asignar a una compañía y cambiar el papel",
    "Asignar una persona a una compañía, cambiarle el papel y quitar la asignación.",
    "Las tres funcionan. El papel va en la asignación, no en la persona.")
add(M, "/admin/personas", "admin@iwl.test", "Corregir un correo mal escrito",
    "Corregir el correo de una cuenta.",
    "Se corrige y esa persona entra con el nuevo.")
add(M, "/admin/personas", "admin@iwl.test", "Archivar y borrar una persona",
    "Archivar una cuenta; borrar otra que no tenga nada registrado.",
    "Archivar le quita el acceso. Borrar solo deja si no arrastra histórico.")
add(M, "/admin/evaluacion", "admin@iwl.test", "Pesos de dimensión y de área",
    "Cambiar un peso y volver a un score.",
    "El score se recalcula. Peso 0 desactiva la dimensión.")
add(M, "/admin/evaluacion", "admin@iwl.test", "Objetivos por etapa",
    "Cambiar el objetivo de una dimensión para una etapa.",
    "La comparativa pasa a medir contra el nuevo objetivo.")
add(M, "/admin/programa", "admin@iwl.test", "Umbrales, bandas y objetivos de tracción",
    "Tocar los umbrales de preparación y las bandas de la cohorte.",
    "Los semáforos y la lectura de la cohorte cambian en consecuencia.")
add(M, "/admin/programa", "admin@iwl.test", "Tarifas por perfil",
    "Fijar la tarifa aplicada y la de mercado de un perfil.",
    "Se guardan y se usan en las horas NUEVAS.")
add(M, "/admin/companias", "admin@iwl.test", "Crear una cohorte",
    "Crear una cohorte nueva.",
    "Queda disponible para dar de alta compañías y abrir convocatoria.")
add(M, "/admin/companias", "admin@iwl.test", "Objetivo interno de cada compañía",
    "Fijar el objetivo interno de una compañía.",
    "Se guarda y se ve en la lectura de la cohorte.")

# --- Informes
M = "20. Informes"
add(M, "/informe/tecnico-inversor/<slug>", "programa@iwl.test", "Informe técnico para inversor",
    "Abrir el informe e imprimirlo a PDF.",
    "Sale completo, con los datos de la compañía y sin cortes de página feos.")
add(M, "/informe/mensual/<slug>", "programa@iwl.test", "Informe mensual",
    "Abrir el informe mensual.",
    "Trae KPI, update y avance del mes.")
add(M, "/informe/aportacion/<slug>", "programa@iwl.test", "Informe de aportación",
    "Abrir el informe de aportación.",
    "Horas, valor y quién aportó. Cuadra con el extracto.")
add(M, "Terminal", "—", "Los cuatro informes en PDF de una vez",
    "Ejecutar: npm run informes",
    "Genera los cuatro PDF de la compañía sin error.")

# --- Permisos
M = "21. Permisos"
add(M, "Varias", "fundadora@marea.test", "Una fundadora solo ve lo suyo",
    "Probar a abrir el proyecto de otra compañía por su dirección.",
    "No entra. Ni lee ni escribe.")
add(M, "Varias", "cto@marea.test", "Segunda persona del mismo equipo",
    "Entrar como la segunda cuenta de Marea y recorrer el proyecto.",
    "Ve lo mismo que la fundadora de esa compañía.")
add(M, "Varias", "mentor@iwl.test", "El mentor que coordina frente al que apoya",
    "Entrar como mentor@iwl.test en Marea (coordina) y en Raíz (apoya).",
    "Puede más en la que coordina. La diferencia está documentada a propósito.")
add(M, "Varias", "programa@iwl.test", "El equipo de IWL ve toda la cartera",
    "Recorrer varias compañías como equipo_iwl.",
    "Las ve todas, pero no entra en /admin.")
add(M, "Varias", "revisor@niage.test", "El revisor ve solo las que lleva",
    "Probar las tres compañías como revisor@niage.test.",
    "Marea y Raíz sí; Vega no.")

# --- Interfaz
M = "22. Interfaz y uso"
add(M, "Todas", "—", "El menú lateral",
    "Recorrer el menú lateral entrando y saliendo de secciones.",
    "Marca dónde estás y lleva a donde dice.")
add(M, "Todas", "—", "Papel claro y lateral oscuro",
    "Comparar la vista de la compañía y la de IWL.",
    "El área de trabajo va en claro en las dos. Lo oscuro es el lateral y la entrada.")
add(M, "Todas", "—", "Varios anchos de ventana",
    "Estrechar la ventana poco a poco: 1512, 1280, 1180, 1024, 820, 600, 390.",
    "Nada se solapa ni se corta, y no hay que desplazar en horizontal.")
add(M, "Todas", "—", "En el móvil",
    "Abrir la plataforma en el teléfono y recorrer tres o cuatro pantallas.",
    "Se usa de verdad: se lee, se pulsa y los formularios se rellenan.")
add(M, "Todas", "—", "Qué se puede pulsar se ve SIN pasar el ratón",
    "Mirar una pantalla quieto, sin mover el ratón. Y probarlo en el móvil, donde no hay ratón.",
    "Las pestañas de sección y las acciones (Editar, Archivar…) llevan cápsula con borde. No hace falta pasar por encima.")
add(M, "Todas", "—", "Los iconos de sección llevan color",
    "Mirar la fila de pestañas de una compañía y la de administración.",
    "Cada icono con su tono, y el de la sección abierta en magenta. El rótulo va siempre al lado.")
add(M, "Todas", "—", "El filete amarillo de los bloques",
    "Mirar el título de cualquier bloque.",
    "Lleva un pelo amarillo debajo. No aparece en ningún gráfico: ahí manda un solo color.")
add(M, "Todas", "—", "No queda «graduación» por ninguna parte",
    "Buscar la palabra en fases, motivos de archivo y desplegables.",
    "Ahora pone «Fin del programa».")
add(M, "Todas", "—", "Moverse solo con el teclado",
    "Recorrer una pantalla con el tabulador.",
    "Siempre se ve dónde está el foco. Nada queda inalcanzable.")
add(M, "Todas", "—", "Los gráficos no dependen del color",
    "Mirar los gráficos de cartera, comparativa y KPI.",
    "Cada serie lleva etiqueta de texto. Dos series parecidas van en dos paneles o con relleno distinto.")
add(M, "Todas", "—", "Los estados vacíos dicen qué hacer",
    "Buscar pantallas sin datos todavía.",
    "Dicen qué hacer para que haya algo, no «no hay nada».")
add(M, "Todas", "—", "Los errores están escritos para una persona",
    "Provocar dos o tres errores a propósito.",
    "Dicen qué ha pasado y qué hacer. Sin signos de exclamación y sin jerga.")
add(M, "Todas", "—", "Un formulario que falla conserva lo escrito",
    "En CUALQUIER formulario largo, forzar un error de validación.",
    "Lo escrito sigue ahí. Solo hay que volver a elegir los ficheros, que el navegador no deja reponer.")
add(M, "Todas", "—", "Una confirmación no desaparece sola",
    "Hacer acciones que cierran su propio panel (logo, firmar, dar cuenta).",
    "La confirmación se lee. Las que llevan contraseña, con más razón.")

# ---------------------------------------------------------------- escribir
cab_fill = PatternFill("solid", fgColor=CARBON)
cab_font = Font(color="FFFFFF", bold=True, size=11)
mod_font = Font(bold=True, color=ZINC)
borde = Border(bottom=Side(style="thin", color=FILETE),
               right=Side(style="thin", color=FILETE))

for i, (titulo, ancho) in enumerate(CABECERAS, start=1):
    c = ws.cell(row=1, column=i, value=titulo)
    c.fill = cab_fill
    c.font = cab_font
    c.alignment = Alignment(vertical="center", horizontal="left")
    ws.column_dimensions[get_column_letter(i)].width = ancho
ws.row_dimensions[1].height = 26

prefijos = {}
fila = 2
for mod, pant, rol, que, pasos, esperado in P:
    n = mod.split(".")[0]
    prefijos[mod] = prefijos.get(mod, 0) + 1
    ident = "%s.%02d" % (n, prefijos[mod])
    valores = [ident, mod, pant, rol, que, pasos, esperado, "", "", ""]
    for col, v in enumerate(valores, start=1):
        c = ws.cell(row=fila, column=col, value=v)
        c.alignment = Alignment(vertical="top", wrap_text=(col >= 5))
        c.border = borde
        if col == 2:
            c.font = mod_font
        if col == 1:
            c.font = Font(color=ZINC)
    ws.row_dimensions[fila].height = 42
    fila += 1

ultima = fila - 1

# Desplegables
dv_res = DataValidation(type="list", formula1='"OK,Falla,A medias,Bloqueado,No aplica"',
                        allow_blank=True, showDropDown=False)
dv_res.error = "Elige uno de la lista."
dv_res.promptTitle = "Resultado"
dv_res.prompt = "OK · Falla · A medias · Bloqueado · No aplica"
ws.add_data_validation(dv_res)
dv_res.add("H2:H%d" % ultima)

dv_gra = DataValidation(type="list", formula1='"Bloqueante,Importante,Menor,Cosmético"',
                        allow_blank=True, showDropDown=False)
dv_gra.promptTitle = "Gravedad"
dv_gra.prompt = "Solo si falla."
ws.add_data_validation(dv_gra)
dv_gra.add("I2:I%d" % ultima)

# Colores según resultado
ws.conditional_formatting.add("H2:H%d" % ultima,
    CellIsRule(operator="equal", formula=['"OK"'],
               fill=PatternFill("solid", fgColor="DCFCE7"), font=Font(color="166534", bold=True)))
ws.conditional_formatting.add("H2:H%d" % ultima,
    CellIsRule(operator="equal", formula=['"Falla"'],
               fill=PatternFill("solid", fgColor="FEE2E2"), font=Font(color="991B1B", bold=True)))
ws.conditional_formatting.add("H2:H%d" % ultima,
    CellIsRule(operator="equal", formula=['"A medias"'],
               fill=PatternFill("solid", fgColor="FEF3C7"), font=Font(color="92400E", bold=True)))
ws.conditional_formatting.add("H2:H%d" % ultima,
    CellIsRule(operator="equal", formula=['"Bloqueado"'],
               fill=PatternFill("solid", fgColor="E4E4E7"), font=Font(color="3F3F46", bold=True)))
ws.conditional_formatting.add("I2:I%d" % ultima,
    CellIsRule(operator="equal", formula=['"Bloqueante"'],
               fill=PatternFill("solid", fgColor="FFE4F0"), font=Font(color="9D0050", bold=True)))

ws.auto_filter.ref = "A1:J%d" % ultima
ws.freeze_panes = "E2"

print("filas", ultima - 1)

# ---------------------------------------------------------------- Cómo se usa
guia = wb.create_sheet("Cómo se usa", 0)
guia.column_dimensions["A"].width = 30
guia.column_dimensions["B"].width = 34
guia.column_dimensions["C"].width = 62

def linea(fila, a, b="", c="", titulo=False, sub=False):
    ca = guia.cell(row=fila, column=1, value=a)
    cb = guia.cell(row=fila, column=2, value=b)
    cc = guia.cell(row=fila, column=3, value=c)
    for cel in (ca, cb, cc):
        cel.alignment = Alignment(vertical="top", wrap_text=True)
    if titulo:
        ca.font = Font(bold=True, size=16, color=MAGENTA)
    elif sub:
        ca.font = Font(bold=True, size=12, color=CARBON)
        guia.row_dimensions[fila].height = 24
    return fila + 1

f = 1
f = linea(f, "Plataforma IWL · pruebas", "", "", titulo=True)
f = linea(f, "", "", "Recorrido de todos los módulos. Una fila por cosa que probar.")
f += 1
f = linea(f, "Cómo se rellena", sub=True)
f = linea(f, "Resultado", "OK / Falla / A medias / Bloqueado / No aplica",
          "Es un desplegable: se pinta solo del color que toca.")
f = linea(f, "Gravedad", "Bloqueante / Importante / Menor / Cosmético",
          "Solo si falla. Bloqueante es lo que impide seguir probando o no se puede enseñar a nadie.")
f = linea(f, "Comentarios", "Texto libre",
          "Qué viste exactamente. Si puedes, con qué datos y en qué pantalla. Una captura vale por tres frases.")
f += 1
f = linea(f, "Para moverse", sub=True)
f = linea(f, "Filtros", "Fila 1", "Cada columna filtra. Para hacer un módulo entero, filtra por Módulo.")
f = linea(f, "Columnas congeladas", "Nº a Entra como", "Se quedan fijas al desplazar a la derecha.")
f = linea(f, "Resumen", "Pestaña «Resumen»", "Cuenta sola según vas rellenando. No hay que tocar nada.")
f += 1
f = linea(f, "Dónde probar", sub=True)
f = linea(f, "En la nube", "https://iwl-platform.vercel.app", "Lo que de verdad va a usar la gente. Ojo: lo que escribas aquí queda.")
f = linea(f, "En local", "npm run db:reset && npm run dev", "Para romper cosas sin miedo. Se reinicia entero en un minuto.")
f = linea(f, "Correos en local", "http://127.0.0.1:54324", "Los enlaces de entrada no salen a internet: caen en esa bandeja.")
f += 1
f = linea(f, "Con qué cuenta entrar (en local)", sub=True)
f = linea(f, "Contraseña común", "iwl-local-2026", "Solo en local. En la nube cada cuenta tiene la suya, en .accesos-nube.md")
for correo, rol, ve in [
    ("admin@iwl.test", "admin_iwl", "Todo, incluida la configuración del programa"),
    ("programa@iwl.test", "equipo_iwl", "Toda la cartera, pero no la configuración"),
    ("revisor@niage.test", "revisor_niage", "Marea Clínica y Raíz Sensórica"),
    ("revisor2@niage.test", "revisor_niage", "Vega Predictiva"),
    ("mentor@iwl.test", "mentor", "Coordina Marea Clínica, apoya en Raíz Sensórica"),
    ("mentor2@iwl.test", "mentor", "Coordina Raíz Sensórica, apoya en Marea Clínica"),
    ("fundadora@marea.test", "fundadora", "Marea Clínica"),
    ("fundadora@vega.test", "fundadora", "Vega Predictiva"),
    ("fundadora@raiz.test", "fundadora", "Raíz Sensórica"),
    ("cto@marea.test", "fundadora", "Marea Clínica, como segunda del equipo"),
]:
    f = linea(f, correo, rol, ve)
f += 1
f = linea(f, "Por dónde empezar", sub=True)
f = linea(f, "1", "Módulos 2 a 7", "El embudo de punta a punta: convocatoria, candidatura, NDA, acuerdo y firma. Es lo más nuevo y lo menos rodado.")
f = linea(f, "2", "Módulo 21", "Permisos. Entrando como quien MENOS puede: ahí es donde se ven los agujeros.")
f = linea(f, "3", "El resto", "Por orden, o por lo que más te interese enseñar.")
f += 1
f = linea(f, "Dos cosas que conviene saber", sub=True)
f = linea(f, "Lo sin medir no es un cero", "", "Un proyecto sin evaluar tiene que parecer sin evaluar («—»), no malo. Si ves un 0, es un fallo.")
f = linea(f, "Nadie valida su propio trabajo", "", "Una fundadora no puntúa ni valida sus puntos. Si en algún sitio puede, es un fallo grave.")

# ---------------------------------------------------------------- Resumen
res = wb.create_sheet("Resumen")
res.column_dimensions["A"].width = 30
for col in "BCDEFG":
    res.column_dimensions[col].width = 13

res["A1"] = "Cómo va"
res["A1"].font = Font(bold=True, size=16, color=MAGENTA)

enc = ["Módulo", "Pruebas", "OK", "Falla", "A medias", "Bloqueado", "Sin probar"]
for i, t in enumerate(enc, start=1):
    c = res.cell(row=3, column=i, value=t)
    c.fill = cab_fill
    c.font = cab_font
res.row_dimensions[3].height = 24

modulos = list(prefijos.keys())
fr = 4
for mod in modulos:
    res.cell(row=fr, column=1, value=mod).font = Font(bold=True, color=ZINC)
    m = '"%s"' % mod
    res.cell(row=fr, column=2, value="=COUNTIF(Pruebas!$B$2:$B$%d,%s)" % (ultima, m))
    for j, estado in enumerate(["OK", "Falla", "A medias", "Bloqueado"], start=3):
        res.cell(row=fr, column=j,
                 value='=COUNTIFS(Pruebas!$B$2:$B$%d,%s,Pruebas!$H$2:$H$%d,"%s")' % (ultima, m, ultima, estado))
    res.cell(row=fr, column=7,
             value='=COUNTIFS(Pruebas!$B$2:$B$%d,%s,Pruebas!$H$2:$H$%d,"")' % (ultima, m, ultima))
    fr += 1

tot = fr
res.cell(row=tot, column=1, value="Total").font = Font(bold=True, color=CARBON)
for col in range(2, 8):
    L = get_column_letter(col)
    c = res.cell(row=tot, column=col, value="=SUM(%s4:%s%d)" % (L, L, tot - 1))
    c.font = Font(bold=True, color=CARBON)
    c.fill = PatternFill("solid", fgColor="F4F4F5")

res.cell(row=tot + 2, column=1, value="Bloqueantes abiertos").font = Font(bold=True, color=MAGENTA)
res.cell(row=tot + 2, column=2,
         value='=COUNTIF(Pruebas!$I$2:$I$%d,"Bloqueante")' % ultima).font = Font(bold=True, color=MAGENTA)

for r in (4, tot):
    pass
res.conditional_formatting.add("D4:D%d" % tot,
    CellIsRule(operator="greaterThan", formula=["0"],
               fill=PatternFill("solid", fgColor="FEE2E2"), font=Font(color="991B1B", bold=True)))
res.conditional_formatting.add("C4:C%d" % tot,
    CellIsRule(operator="greaterThan", formula=["0"],
               fill=PatternFill("solid", fgColor="DCFCE7"), font=Font(color="166534")))

wb.active = 0
import os
destino = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                       "IWL_Checklist_Pruebas.xlsx")
wb.save(destino)
print("guardado en", destino)
print("pruebas:", ultima - 1, "· modulos:", len(modulos))
